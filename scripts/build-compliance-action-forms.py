from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "docs" / "compliance-action-forms"

FORMS = [
    ("양식 1", "안전 및 유지관리계획서", "계획"),
    ("양식 2", "도급·용역·위탁 계약관리서", "계약"),
    ("양식 3", "정밀안전진단 실시·결과서", "정밀안전진단"),
    ("양식 4", "안전점검 실시·결과서", "안전점검"),
    ("양식 5", "기타 조치 및 증빙자료서", "기타"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_label_row(table, left_label: str, right_label: str) -> None:
    cells = table.add_row().cells
    set_cell_text(cells[0], left_label, True)
    set_cell_shading(cells[0], "EAF1F8")
    set_cell_text(cells[1], "")
    set_cell_text(cells[2], right_label, True)
    set_cell_shading(cells[2], "EAF1F8")
    set_cell_text(cells[3], "")


def build_form(form_no: str, title: str, category: str) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"[{form_no}] {category} 기록 양식")
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(10)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(18)

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    add_label_row(table, "관리대상명", "관련 의무명")
    add_label_row(table, "담당 부서", "책임자")
    add_label_row(table, "수행 업체명", "사업자번호")
    add_label_row(table, "대표자", "연락처")
    add_label_row(table, "조치이행 시작일", "조치이행 종료일")
    add_label_row(table, "등록일", "기록구분")

    document.add_paragraph()
    detail_table = document.add_table(rows=2, cols=1)
    detail_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    detail_table.style = "Table Grid"
    set_cell_text(detail_table.cell(0, 0), "조치내용 요약", True, 11)
    set_cell_shading(detail_table.cell(0, 0), "EAF1F8")
    set_cell_text(detail_table.cell(1, 0), "첨부문서의 핵심 조치내용을 간략히 입력하세요.", False, 10)
    detail_table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    detail_table.cell(1, 0).height = Cm(5.5)

    document.add_paragraph()
    evidence_table = document.add_table(rows=6, cols=3)
    evidence_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    evidence_table.style = "Table Grid"
    headers = ["번호", "첨부문서명", "등록일"]
    for index, header in enumerate(headers):
        set_cell_text(evidence_table.cell(0, index), header, True)
        set_cell_shading(evidence_table.cell(0, index), "EAF1F8")
    for row in range(1, 6):
        set_cell_text(evidence_table.cell(row, 0), str(row))
        set_cell_text(evidence_table.cell(row, 1), "")
        set_cell_text(evidence_table.cell(row, 2), "")

    document.add_paragraph()
    note = document.add_paragraph(
        "※ 조치이행 기간은 실제 업무 수행 기간이며, 등록일은 시스템에 문서와 기록을 등록한 날짜입니다."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.name = "Malgun Gothic"
    note.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"{form_no.replace(' ', '_')}_{title}.docx"
    output_path = OUTPUT_DIR / filename
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    paths = [build_form(*form) for form in FORMS]
    for path in paths:
        print(path)
